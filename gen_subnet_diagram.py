from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE

# ============================================================
# カラー
# ============================================================
C_AZURE        = RGBColor(0x00, 0x78, 0xD4)
C_VNET_BG      = RGBColor(0xE3, 0xF2, 0xFD)
C_SUBNET_INT   = RGBColor(0xB3, 0xD9, 0xF7)  # subnet-integration（青系）
C_SUBNET_PE    = RGBColor(0xB2, 0xDF, 0xDB)  # subnet-pe（緑系）
C_SUBNET_GW    = RGBColor(0xE8, 0xEA, 0xED)  # GatewaySubnet（グレー）
C_ONPREM       = RGBColor(0xF5, 0xF5, 0xF5)
C_APP_SVC      = RGBColor(0xFF, 0xFF, 0xFF)
C_FOUNDRY      = RGBColor(0xF3, 0xE5, 0xF5)
C_PE           = RGBColor(0xE0, 0xF7, 0xFA)
C_MANAGED      = RGBColor(0xE8, 0xF5, 0xE9)
C_ARROW        = RGBColor(0x01, 0x57, 0x9B)
C_ARROW_OUT    = RGBColor(0x1B, 0x5E, 0x20)
C_WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
C_DARK         = RGBColor(0x1A, 0x1A, 0x1A)
C_GRAY         = RGBColor(0x60, 0x60, 0x60)
C_RED          = RGBColor(0xC6, 0x28, 0x28)
C_GREEN        = RGBColor(0x1B, 0x5E, 0x20)
C_BORDER_INT   = RGBColor(0x01, 0x57, 0x9B)
C_BORDER_PE    = RGBColor(0x00, 0x69, 0x5C)
C_BORDER_GW    = RGBColor(0x90, 0x90, 0x90)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

# ============================================================
# ヘルパー
# ============================================================
def rect(slide, x, y, w, h, fill, border=None, bw=1.5):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if border:
        s.line.color.rgb = border; s.line.width = Pt(bw)
    else:
        s.line.fill.background()
    return s

def txt(slide, x, y, w, h, text, size=10, bold=False,
        color=None, align=PP_ALIGN.CENTER, wrap=True):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.word_wrap = wrap
    tf = t.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color or C_DARK
    return t

def box(slide, x, y, w, h, label, sub=None,
        fill=C_WHITE, border=C_AZURE, lsize=9, ssize=7.5):
    rect(slide, x, y, w, h, fill, border, 1.2)
    txt(slide, x, y+0.04, w, 0.22, label, lsize, True, C_DARK)
    if sub:
        txt(slide, x, y+0.25, w, 0.2, sub, ssize, False, C_GRAY)
    return

def arrow(slide, x1, y1, x2, y2, color=C_ARROW, w=1.8):
    c = slide.shapes.add_connector(2,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)

def label_arrow(slide, x, y, w, h, text, color=C_ARROW):
    txt(slide, x, y, w, h, text, 7.5, False, color)

# ============================================================
# スライド 1: App Service と Subnet の関係（概念）
# ============================================================
s1 = prs.slides.add_slide(prs.slide_layouts[6])

txt(s1, 0.2, 0.08, 12.9, 0.38,
    "App Service と Subnet の関係 — 概念図",
    15, True, C_AZURE)

# --- 説明テキスト ---
txt(s1, 0.3, 0.55, 12.5, 0.3,
    "App Service は Subnet の「中」に置くのではなく、VNet の外側（Azure マネージドゾーン）に存在します。Subnet との関係は 2 種類あります。",
    8.5, False, C_GRAY)

# ======= 左カラム: VNet Integration =======
rect(s1, 0.3, 1.05, 5.8, 5.6, C_VNET_BG, C_AZURE, 2)
txt(s1, 0.3, 1.07, 5.8, 0.28, "パターン①  VNet Integration（アウトバウンド）", 9, True, C_AZURE)

# Subnet-integration
rect(s1, 0.55, 1.5, 3.2, 1.5, C_SUBNET_INT, C_BORDER_INT, 1.5)
txt(s1, 0.55, 1.52, 3.2, 0.25, "subnet-integration", 8.5, True, C_BORDER_INT)
txt(s1, 0.65, 1.8, 3.0, 0.6,
    "App Service から VNet 内への\n通信が通る「出口」", 8, False, C_DARK)

# App Service box (outside VNet)
rect(s1, 4.2, 1.5, 1.65, 1.5, C_APP_SVC, C_AZURE, 1.5)
txt(s1, 4.2, 1.52, 1.65, 0.25, "App Service", 8.5, True, C_AZURE)
txt(s1, 4.2, 1.8, 1.65, 0.6, "Azure\nマネージド\nゾーン", 7.5, False, C_GRAY)

# 矢印: App Service → subnet-integration
arrow(s1, 4.2, 2.25, 3.75, 2.25, C_ARROW_OUT, 2)
label_arrow(s1, 3.3, 2.05, 1.5, 0.22, "アウトバウンド\n（出る方向）", C_ARROW_OUT)

# 説明
rect(s1, 0.4, 3.2, 5.5, 1.3, RGBColor(0xE3,0xF2,0xFD), C_BORDER_INT, 0.8)
txt(s1, 0.5, 3.28, 5.3, 1.1,
    "用途例:\n• Backend → Foundry への通信\n• Backend → Azure Functions（Tool/MCP）への通信\n• Frontend → Backend への通信",
    8, False, C_DARK, PP_ALIGN.LEFT)

txt(s1, 0.4, 4.6, 5.5, 0.9,
    "⚠  委任（Delegation）が必要\n同一 App Service Plan なら Subnet を共用できる\n別プランなら別 Subnet が必要",
    8, False, C_RED, PP_ALIGN.LEFT)

# ======= 右カラム: Private Endpoint =======
rect(s1, 6.8, 1.05, 6.1, 5.6, C_VNET_BG, C_AZURE, 2)
txt(s1, 6.8, 1.07, 6.1, 0.28, "パターン②  Private Endpoint（インバウンド）", 9, True, C_AZURE)

# subnet-pe
rect(s1, 7.05, 1.5, 5.6, 3.8, C_SUBNET_PE, C_BORDER_PE, 1.5)
txt(s1, 7.05, 1.52, 5.6, 0.25, "subnet-pe", 8.5, True, C_BORDER_PE)

# Private Endpoint（仮想NIC）
rect(s1, 7.25, 1.85, 2.0, 1.2, C_PE, C_BORDER_PE, 1.0)
txt(s1, 7.25, 1.87, 2.0, 0.25, "Private Endpoint", 8, True, C_BORDER_PE)
txt(s1, 7.25, 2.13, 2.0, 0.7,
    "Subnet 内に置く\n仮想 NIC\n（プライベート IP）", 7.5, False, C_DARK)

# App Service box
rect(s1, 10.0, 1.85, 1.6, 1.2, C_APP_SVC, C_AZURE, 1.5)
txt(s1, 10.0, 1.87, 1.6, 0.25, "App Service", 8.5, True, C_AZURE)
txt(s1, 10.0, 2.13, 1.6, 0.7, "Azure\nマネージドゾーン", 7.5, False, C_GRAY)

# 矢印: PE → App Service
arrow(s1, 9.25, 2.45, 10.0, 2.45, C_BORDER_PE, 2)
label_arrow(s1, 9.15, 2.25, 1.1, 0.22, "プライベート\n接続", C_BORDER_PE)

# Public 無効化バッジ
rect(s1, 10.0, 3.3, 1.6, 0.35, C_RED, None)
txt(s1, 10.0, 3.32, 1.6, 0.28, "Public 無効", 7.5, True, C_WHITE)

# 説明
rect(s1, 7.15, 3.8, 5.5, 1.3, RGBColor(0xE0,0xF7,0xEF), C_BORDER_PE, 0.8)
txt(s1, 7.25, 3.88, 5.3, 1.1,
    "用途例:\n• VPN 経由ユーザー → Frontend\n• Frontend → Backend（VNet 内通信）\n• Backend → Foundry / Functions",
    8, False, C_DARK, PP_ALIGN.LEFT)

txt(s1, 7.15, 5.2, 5.5, 0.35,
    "Public Endpoint を無効化 → 社外から存在が見えない",
    8, False, C_RED, PP_ALIGN.LEFT)

# 重要: 2 つを同一 Subnetに置けない
rect(s1, 0.3, 6.35, 12.7, 0.55, RGBColor(0xFF,0xF3,0xE0), RGBColor(0xE6,0x51,0x00), 1.2)
txt(s1, 0.4, 0.38+6.0, 12.5, 0.45,
    "⚠  重要: VNet Integration（①）と Private Endpoint（②）は同一 Subnet に共存できません → 必ず別 Subnet に分ける",
    9, True, RGBColor(0xE6, 0x51, 0x00))

# ============================================================
# スライド 2: 全体構成図（従属関係）
# ============================================================
s2 = prs.slides.add_slide(prs.slide_layouts[6])

txt(s2, 0.2, 0.08, 12.9, 0.38,
    "全体構成図 — リソースと Subnet の従属関係",
    15, True, C_AZURE)

# ======= オンプレ =======
rect(s2, 0.15, 0.55, 1.9, 2.1, C_ONPREM, C_GRAY, 1.2)
txt(s2, 0.15, 0.57, 1.9, 0.25, "オンプレミス", 8, True, C_GRAY)
box(s2, 0.25, 0.88, 1.65, 0.62, "ユーザー PC", "ブラウザ",
    fill=C_WHITE, border=C_GRAY)
box(s2, 0.25, 1.6, 1.65, 0.62, "VPN デバイス", "IPsec",
    fill=C_WHITE, border=C_GRAY)

# ======= VNet 大枠 =======
rect(s2, 2.3, 0.48, 7.6, 6.7, C_VNET_BG, C_AZURE, 2.2)
txt(s2, 2.3, 0.5, 7.6, 0.3, "Virtual Network  10.0.0.0/16", 9, True, C_AZURE)

# --- GatewaySubnet ---
rect(s2, 2.5, 0.88, 2.5, 1.0, C_SUBNET_GW, C_BORDER_GW, 1.2)
txt(s2, 2.5, 0.9, 2.5, 0.25, "GatewaySubnet  10.0.255.0/27", 7.5, True, C_GRAY)
box(s2, 2.6, 1.2, 2.3, 0.55, "VPN Gateway", "pip: output で確認",
    fill=C_WHITE, border=C_GRAY)

# --- subnet-pe ---
rect(s2, 2.5, 2.05, 2.5, 4.85, C_SUBNET_PE, C_BORDER_PE, 1.5)
txt(s2, 2.5, 2.07, 2.5, 0.25, "subnet-pe  10.0.2.0/24", 7.5, True, C_BORDER_PE)

box(s2, 2.6, 2.42, 2.25, 0.62, "PE: Frontend", "Private Endpoint",
    fill=C_PE, border=C_BORDER_PE)
box(s2, 2.6, 3.18, 2.25, 0.62, "PE: Backend", "Private Endpoint",
    fill=C_PE, border=C_BORDER_PE)
box(s2, 2.6, 3.94, 2.25, 0.62, "PE: Functions", "Private Endpoint",
    fill=C_PE, border=C_BORDER_PE)
box(s2, 2.6, 4.7, 2.25, 0.62, "PE: Foundry", "Private Endpoint",
    fill=C_PE, border=C_BORDER_PE)
box(s2, 2.6, 5.46, 2.25, 0.45, "Private DNS Zone", "privatelink.*",
    fill=C_MANAGED, border=C_GREEN, lsize=7.5)

# --- subnet-integration ---
rect(s2, 5.25, 2.05, 4.4, 4.85, C_SUBNET_INT, C_BORDER_INT, 1.5)
txt(s2, 5.25, 2.07, 4.4, 0.25, "subnet-integration  10.0.1.0/24  （VNet Integration 委任済み）",
    7.5, True, C_BORDER_INT)

# App Service Plan バッジ
rect(s2, 5.4, 2.42, 4.1, 0.25, RGBColor(0xE3,0xF2,0xFD), C_AZURE, 0.8)
txt(s2, 5.4, 2.43, 4.1, 0.22, "App Service Plan（B1 / Dedicated）— 3 リソース共用",
    7, True, C_AZURE)

box(s2, 5.4, 2.78, 1.85, 0.75, "App Service", "Frontend\nassistant-ui",
    fill=C_APP_SVC, border=C_AZURE)
box(s2, 7.4, 2.78, 1.85, 0.75, "App Service", "Backend\nFastAPI",
    fill=C_APP_SVC, border=C_AZURE)
box(s2, 5.4, 3.68, 3.85, 0.75, "Azure Functions", "Tool / MCP サーバー",
    fill=C_APP_SVC, border=C_AZURE)

# Managed ID バッジ
rect(s2, 5.4, 4.55, 1.85, 0.22, C_MANAGED, C_GREEN, 0.8)
txt(s2, 5.4, 4.56, 1.85, 0.2, "Managed ID", 7, True, C_GREEN)
rect(s2, 7.4, 4.55, 1.85, 0.22, C_MANAGED, C_GREEN, 0.8)
txt(s2, 7.4, 4.56, 1.85, 0.2, "Managed ID", 7, True, C_GREEN)

txt(s2, 5.4, 4.85, 3.85, 0.8,
    "VNet Integration とは\nApp Service がこの Subnet を「出口」として使い\nVNet 内の他リソースと通信できる仕組み",
    7.5, False, C_DARK, PP_ALIGN.LEFT)

# ======= Azure マネージドゾーン =======
rect(s2, 10.15, 0.48, 2.95, 6.7, C_FOUNDRY, RGBColor(0x7B, 0x1F, 0xA2), 1.5)
txt(s2, 10.15, 0.5, 2.95, 0.3, "Microsoft Foundry", 9, True, RGBColor(0x6A,0x1A,0x9A))

box(s2, 10.3, 0.9, 2.6, 0.7, "親エージェント", "Foundry Agent",
    fill=C_WHITE, border=RGBColor(0x7B,0x1F,0xA2))
box(s2, 10.3, 1.75, 2.6, 0.7, "子エージェント", "",
    fill=C_WHITE, border=RGBColor(0x7B,0x1F,0xA2))
box(s2, 10.3, 2.6, 2.6, 0.7, "GPT-5.2", "LLM",
    fill=C_WHITE, border=RGBColor(0x7B,0x1F,0xA2))
box(s2, 10.3, 3.45, 2.6, 0.7, "Embedding 003", "埋め込み",
    fill=C_WHITE, border=RGBColor(0x7B,0x1F,0xA2))

# ======= 矢印群 =======
# オンプレ → GatewaySubnet
arrow(s2, 2.05, 1.91, 2.5, 1.91)
label_arrow(s2, 1.82, 1.72, 0.9, 0.2, "VPN", C_ARROW)

# VPN GW → subnet-pe （下方向）
arrow(s2, 3.75, 1.75, 3.75, 2.42)

# PE: Frontend → App Service Frontend
arrow(s2, 4.85, 2.73, 5.4, 2.73, C_BORDER_PE, 1.5)
# PE: Backend → App Service Backend
arrow(s2, 4.85, 3.49, 7.4, 3.16, C_BORDER_PE, 1.5)
# PE: Functions → Azure Functions
arrow(s2, 4.85, 4.25, 5.4, 4.05, C_BORDER_PE, 1.5)
# PE: Foundry → Foundry
arrow(s2, 4.85, 5.01, 10.15, 1.6, C_BORDER_PE, 1.5)

# Frontend → PE:Backend （VNet内通信）
arrow(s2, 6.33, 3.53, 4.85, 3.49, C_ARROW_OUT, 1.5)
label_arrow(s2, 5.3, 3.57, 1.3, 0.2, "VNet 内通信", C_ARROW_OUT)

# Backend → PE:Functions
arrow(s2, 7.4, 4.16, 4.85, 4.25, C_ARROW_OUT, 1.5)
# Backend → PE:Foundry
arrow(s2, 8.33, 3.53, 8.33, 5.3, C_ARROW_OUT, 1.5)
arrow(s2, 8.33, 5.3, 4.85, 5.01, C_ARROW_OUT, 1.5)

# ============================================================
# スライド 3: 通信フロー（ステップ番号付き）
# ============================================================
s3 = prs.slides.add_slide(prs.slide_layouts[6])

txt(s3, 0.2, 0.08, 12.9, 0.38,
    "通信フロー（ステップ順）",
    15, True, C_AZURE)

steps = [
    ("①", "ユーザー PC → VPN デバイス",
     "オンプレ内通信",
     C_GRAY, C_ONPREM),
    ("②", "VPN デバイス → VPN Gateway",
     "IPsec トンネル（インターネット経由だが暗号化）",
     C_GRAY, C_ONPREM),
    ("③", "VPN Gateway → PE: Frontend  （subnet-pe 内）",
     "VNet に入った通信が subnet-pe の Private Endpoint（仮想NIC）に届く",
     C_BORDER_PE, C_SUBNET_PE),
    ("④", "PE: Frontend → App Service Frontend  （Azure マネージドゾーン）",
     "Private Endpoint がプライベートIPでApp Serviceに転送。Public Endpoint は無効なのでここ以外からは届かない",
     C_AZURE, C_APP_SVC),
    ("⑤", "App Service Frontend → subnet-integration → PE: Backend",
     "VNet Integration（subnet-integration が出口）→ subnet-pe の PE: Backend に到達",
     C_BORDER_INT, C_SUBNET_INT),
    ("⑥", "PE: Backend → App Service Backend",
     "④ と同様",
     C_AZURE, C_APP_SVC),
    ("⑦", "App Service Backend → subnet-integration → PE: Foundry / PE: Functions",
     "Managed ID 認証。VNet Integration 経由で Foundry または Functions の PE に到達",
     C_BORDER_INT, C_SUBNET_INT),
]

for i, (num, title, desc, tcol, bgcol) in enumerate(steps):
    y = 0.55 + i * 0.84
    rect(s3, 0.3, y, 12.7, 0.78, bgcol,
         tcol if bgcol != C_APP_SVC else C_AZURE, 1.0)
    txt(s3, 0.35, y+0.04, 0.45, 0.35, num, 14, True, tcol)
    txt(s3, 0.85, y+0.04, 11.9, 0.28, title, 10, True, C_DARK, PP_ALIGN.LEFT)
    txt(s3, 0.85, y+0.35, 11.9, 0.35, desc, 8.5, False, C_GRAY, PP_ALIGN.LEFT)

# ============================================================
# 保存
# ============================================================
out = "/home/user/tmp/azure-test-environment/docs/architecture.pptx"
prs.save(out)
print(f"Saved: {out}")
